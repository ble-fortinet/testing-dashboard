import io
from datetime import date, datetime

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import PatternFill

from tdash import workbook as wbk
from tdash.workbook import WorkbookError, diff, export, lines, next_id, parse, roster_csv

P_HEAD = ["USERLOGIN", "in_roster", "purpose", "COUNTRY", "HIREDATE"]
C_HEAD = [
    "id",
    "status",
    "category",
    "tests",
    "persona",
    "question",
    "expected",
    "must_not",
    "sources",
    "threshold",
    "notes",
]


def book(cases: list[list], personas: list[list] | None = None, c_head=C_HEAD, p_head=P_HEAD) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Cases"
    ws.append(c_head)
    for r in cases:
        ws.append(r)
    wp = wb.create_sheet("Personas")
    wp.append(p_head)
    for r in personas if personas is not None else [["test-ca-on", "Y", "", "CA", "2021-03-15"]]:
        wp.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def row(**kw) -> list:
    base = {"id": "TC-001", "status": "ready", "category": "Time off", "tests": "", "persona": "test-ca-on", "question": "Q?",
                "expected": "One\nTwo", "must_not": "", "sources": "cite", "threshold": None, "notes": ""}  # fmt: skip
    base.update(kw)
    return [base[k] for k in C_HEAD]


# ---------------------------------------------------------------- the seed (frozen copy in test/fixtures)
def test_seed_parses_cleanly(seed):
    assert len(seed["cases"]) == 75
    assert len(seed["personas"]) == 12
    assert [c for c in seed["cases"] if c["issues"]] == []  # every row is valid
    assert seed["warnings"] == []
    assert all(c["persona"].startswith("test-") for c in seed["cases"])


def test_seed_review_markers_survive(seed):
    kinds = [c["review"]["kind"] for c in seed["cases"] if c["review"]]
    assert kinds.count("add") == 10 and kinds.count("check") == 20
    adds = [c for c in seed["cases"] if c["review"] and c["review"]["kind"] == "add"]
    assert all(c["status"] == "draft" for c in adds)  # the file says so; nothing invented
    first = next(c for c in seed["cases"] if c["review"])
    assert not first["review"]["note"].upper().startswith(("ADD", "CHECK"))  # prefix stripped
    assert first["review"]["col"] == "expected"


def test_seed_persona_comments_and_columns(seed):
    newhire = next(p for p in seed["personas"] if p["USERLOGIN"] == "test-ca-newhire")
    assert newhire["HIREDATE"] == "-10d"
    assert "W1" in newhire["_comments"]["HIREDATE"]
    assert "ACTION" not in seed["persona_columns"] and "needs_input" not in seed["persona_columns"]


def test_seed_round_trips_through_export(seed):
    out = parse(export(seed["cases"], seed["personas"], seed["persona_columns"], seed["case_extra_columns"]))
    snap = {
        "cases": seed["cases"],
        "personas": seed["personas"],
        "persona_columns": seed["persona_columns"],
        "edited": [],
    }
    d = diff(snap, out)
    assert d["cases"]["added"] == d["cases"]["removed"] == [] and d["cases"]["changed"] == []
    assert d["personas"] == {"added": [], "removed": [], "changed": [], "columns_changed": False}


def test_export_keeps_dropdowns_in_a_form_openpyxl_can_read(seed):
    data = export(seed["cases"], seed["personas"], seed["persona_columns"], [])
    wb = load_workbook(io.BytesIO(data))
    assert "PersonaList" in wb.defined_names
    dvs = {str(dv.sqref): dv.formula1 for dv in wb["Cases"].data_validations.dataValidation}
    assert dvs["E2:E500"] == "=PersonaList"  # persona, through the named range
    assert dvs["B2:B500"] == '"ready,draft,off"'
    # a second load/save cycle keeps it (the bug that dropped it once)
    buf = io.BytesIO()
    wb.save(buf)
    again = load_workbook(io.BytesIO(buf.getvalue()))
    assert "=PersonaList" in [dv.formula1 for dv in again["Cases"].data_validations.dataValidation]


def test_export_writes_review_fill_and_comment(seed):
    data = export(seed["cases"], seed["personas"], seed["persona_columns"], [])
    ws = load_workbook(io.BytesIO(data))["Cases"]
    orange = [c for c in ws["G"][1:] if c.fill.fgColor.rgb and c.fill.fgColor.rgb.endswith("F8CBAD")]
    assert len(orange) == 10
    assert all(c.comment.text.startswith("ADD — ") for c in orange)


# ---------------------------------------------------------------- whole-file errors
@pytest.mark.parametrize(
    ("data", "msg"),
    [
        (b"not a zip", "not a readable"),
        (book([row()], c_head=["id", "question", "persona"]), "missing column"),
        (book([row(), row()]), "repeats id TC-001"),
        (book([row()], [["test-a", "Y", "", "CA", ""], ["test-a", "Y", "", "CA", ""]]), "repeats USERLOGIN"),
        (book([row()], [["BLE", "Y", "", "CA", ""]]), "must be test- logins"),
        (book([row()], p_head=["LOGIN", "COUNTRY"]), "no USERLOGIN"),
    ],
)
def test_file_level_errors(data, msg):
    with pytest.raises(WorkbookError, match=msg):
        parse(data)


def test_missing_sheets():
    wb = Workbook()
    wb.active.title = "Cases"
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(WorkbookError, match="Personas"):
        parse(buf.getvalue())


# ---------------------------------------------------------------- row-level issues are kept, not dropped
def test_bad_rows_are_kept_with_issues():
    p = [["test-ca-on", "Y", "", "CA", ""], ["test-us", "Y", "", "US", ""]]
    out = parse(
        book(
            [
                row(id="TC-001", persona="test-nobody"),
                row(id="TC-002", status="maybe"),
                row(id="TC-003", threshold="abc"),
                row(id="TC-004", expected="", must_not="", sources=""),
                row(id="TC-005", sources="always"),
                row(id="TC-006", question=""),
                row(id="TC-007"),
            ],
            p,
        )
    )
    issues = {c["id"]: c["issues"] for c in out["cases"]}
    assert "not on the Personas sheet" in issues["TC-001"][0]
    assert "status must be" in issues["TC-002"][0]
    assert "not a number" in issues["TC-003"][0]
    assert "nothing to grade" in issues["TC-004"][0]
    assert "sources must be" in issues["TC-005"][0]
    assert "question is empty" in issues["TC-006"][0]
    assert issues["TC-007"] == []
    assert len(out["cases"]) == 7


def test_a_case_naming_a_real_login_is_flagged():
    # the Personas sheet itself refuses real logins; a case can still name one
    out = parse(book([row(persona="BLE")]))
    assert "must be a test- login" in out["cases"][0]["issues"][0]


def test_values_are_normalised():
    out = parse(
        book([row(id=None, status="", threshold="80%", expected="- One\n• Two\n3) Three\n\n", must_not="1. Never")])
    )
    c = out["cases"][0]
    assert c["id"] == "TC-001" and "assigned TC-001" in out["warnings"][0]
    assert c["status"] == "ready"
    assert c["threshold"] == 0.8
    assert lines(c["expected"]) == ["One", "Two", "Three"]
    assert c["must_not"] == "Never"


def test_blank_status_on_an_add_row_becomes_draft():
    wb = Workbook()
    ws = wb.active
    ws.title = "Cases"
    ws.append(C_HEAD)
    ws.append(row(status=None))
    ws["G2"].fill = PatternFill("solid", fgColor="F8CBAD")
    wb.create_sheet("Personas").append(P_HEAD)
    wb["Personas"].append(["test-ca-on", "Y", "", "CA", ""])
    buf = io.BytesIO()
    wb.save(buf)
    c = parse(buf.getvalue())["cases"][0]
    assert c["review"] == {"kind": "add", "note": "", "col": "expected"}  # a bare fill still counts
    assert c["status"] == "draft"


def test_comment_cleanup_threaded_and_author():
    wb = Workbook()
    ws = wb.active
    ws.title = "Cases"
    ws.append(C_HEAD)
    ws.append(row(id="TC-001"))
    ws.append(row(id="TC-002"))
    ws.append(row(id="TC-003"))
    threaded = (
        "[Threaded comment]\n\nYour version of Excel allows you to read this threaded comment; however, any edits to it "
        "will get removed if the file is opened in a newer version of Excel. Learn more: https://go.microsoft.com/fwlink/?linkid=870924\n\n"
        "Comment:\n    CHECK — confirm the deadline\nReply:\n    ok"
    )
    ws["G2"].comment = Comment(threaded, "Ben Le")
    ws["G3"].comment = Comment("Ben Le:\nADD: the hotline name", "Ben Le")
    ws["H4"].comment = Comment("just a thought", "HR")  # a comment on another column, no fill
    wb.create_sheet("Personas").append(P_HEAD)
    wb["Personas"].append(["test-ca-on", "Y", "", "CA", ""])
    buf = io.BytesIO()
    wb.save(buf)
    cs = {c["id"]: c for c in parse(buf.getvalue())["cases"]}
    assert cs["TC-001"]["review"] == {"kind": "check", "note": "confirm the deadline", "col": "expected"}
    assert cs["TC-002"]["review"] == {"kind": "add", "note": "the hotline name", "col": "expected"}
    assert cs["TC-003"]["review"] == {"kind": "check", "note": "just a thought", "col": "must_not"}


def test_excel_dates_and_numbers_become_text():
    out = parse(book([row()], [["test-ca-on", "Y", "", "CA", datetime(2021, 3, 15)]]))
    assert out["personas"][0]["HIREDATE"] == "2021-03-15"


def test_extra_columns_are_preserved():
    out = parse(book([[*row(), "HR-owned"]], c_head=[*C_HEAD, "Owner"]))
    assert out["case_extra_columns"] == ["Owner"]
    assert out["cases"][0]["extra"] == {"Owner": "HR-owned"}
    back = parse(export(out["cases"], out["personas"], out["persona_columns"], out["case_extra_columns"]))
    assert back["cases"][0]["extra"] == {"Owner": "HR-owned"}


# ---------------------------------------------------------------- diff
def test_diff_reports_reverted_in_app_edits(seed):
    current = {"cases": seed["cases"], "personas": seed["personas"], "persona_columns": seed["persona_columns"]}
    current["edited"] = ["TC-001", "TC-002"]
    incoming = {**seed, "cases": [dict(c) for c in seed["cases"]]}
    incoming["cases"][0] = {**incoming["cases"][0], "question": "changed?"}
    incoming["cases"] = [c for c in incoming["cases"] if c["id"] != "TC-003"]
    incoming["cases"].append({**seed["cases"][0], "id": "TC-999"})
    d = diff(current, incoming)
    assert d["cases"]["added"] == ["TC-999"]
    assert d["cases"]["removed"] == ["TC-003"]
    assert d["cases"]["changed"] == [{"id": "TC-001", "fields": ["question"]}]
    assert d["cases"]["reverts_in_app_edits"] == ["TC-001"]  # TC-002 is unchanged, so not reverted


# ---------------------------------------------------------------- roster CSV
def test_roster_csv(seed):
    text = roster_csv(seed["personas"], seed["persona_columns"])
    head, *rows = text.strip().split("\n")
    cols = head.split(",")
    assert "USERLOGIN" in cols and "in_roster" not in cols and "purpose" not in cols
    logins = [r.split(",")[0] for r in rows]
    assert "test-unlisted" not in logins and len(logins) == 11
    assert "-10d" in text


def test_roster_csv_resolves_relative_dates(seed):
    text = roster_csv(seed["personas"], seed["persona_columns"], resolve_dates=True, today=date(2026, 10, 9))
    assert "-10d" not in text and "-5d" not in text
    assert "2026-09-29" in text and "2026-10-04" in text


def test_roster_csv_resolves_a_future_date(seed):
    """+Nd is a pre-boarding persona: it starts after today."""
    personas = [*seed["personas"], {"USERLOGIN": "test-pre", "HIREDATE": "+7d"}]
    text = roster_csv(personas, seed["persona_columns"], resolve_dates=True, today=date(2026, 10, 9))
    assert "+7d" not in text and "2026-10-16" in text


def test_persona_dates_accept_both_directions():
    cols = ["USERLOGIN", "HIREDATE"]
    assert wbk.normalize_persona({"USERLOGIN": "test-pre", "HIREDATE": "+7d"}, cols, set())
    assert wbk.normalize_persona({"USERLOGIN": "test-new", "HIREDATE": "-10d"}, cols, set())
    with pytest.raises(ValueError, match="days ahead"):
        wbk.normalize_persona({"USERLOGIN": "test-bad", "HIREDATE": "7d"}, cols, set())


def test_next_id():
    assert next_id({"TC-009", "TC-010", "X"}) == "TC-011"
    assert next_id(set()) == "TC-001"


def test_is_test_login():
    assert wbk.is_test_login("test-ca-on") and wbk.is_test_login("TEST-X")
    assert not wbk.is_test_login("BLE") and not wbk.is_test_login("")


# ---------------------------------------------------------------- export safety
def test_formula_looking_text_is_exported_as_data():
    evil = '=HYPERLINK("http://evil.example","click")'
    out = parse(book([row()]))
    out["cases"][0]["question"] = evil  # as the app would hold it after an in-app edit
    data = export(out["cases"], out["personas"], out["persona_columns"], [])
    ws = load_workbook(io.BytesIO(data))["Cases"]
    assert ws["F2"].data_type == "s" and ws["F2"].value == evil
    assert parse(data)["cases"][0]["question"] == evil  # round-trips as text


def test_control_characters_and_huge_text_dont_break_exports():
    out = parse(book([row()]))
    out["cases"][0]["question"] = "line one\x0bline two\x00"
    data = export(out["cases"], out["personas"], out["persona_columns"], [])
    assert parse(data)["cases"][0]["question"] == "line oneline two"
    out["cases"][0]["notes"] = "=1+1 " + "x" * 40000 + ""
    ws = load_workbook(io.BytesIO(export(out["cases"], out["personas"], out["persona_columns"], [])))["Cases"]
    notes = ws.cell(row=2, column=11).value
    assert notes.startswith("=1+1") and notes.endswith("[truncated]") and len(notes) <= 32767
    assert ws.cell(row=2, column=11).data_type == "s"


@pytest.mark.parametrize("bad", ["TC 001", "TC/001", "-x", "a" * 41])
def test_ids_must_be_url_safe(bad):
    with pytest.raises(WorkbookError, match="letters, digits"):
        parse(book([row(id=bad)]))


# ---------------------------------------------------------------- personas in the app
COLS = wbk.DEFAULT_PERSONA_COLUMNS


def test_a_new_persona_takes_the_sheet_columns():
    p = wbk.normalize_persona({"USERLOGIN": "test-ca-new", "HIREDATE": "-10d", "junk": "x"}, COLS, set())
    assert list(p) == COLS and p["HIREDATE"] == "-10d" and "junk" not in p


@pytest.mark.parametrize(
    ("raw", "why"),
    [
        ({"USERLOGIN": "JSMITH"}, "start test-"),
        ({"USERLOGIN": "test-a b"}, "start test-"),
        ({"USERLOGIN": "test-x", "HIREDATE": "March 1"}, "HIREDATE"),
        ({"USERLOGIN": "test-x", "in_roster": "maybe"}, "in_roster"),
        ({"USERLOGIN": "test-x", "MGRLVL1": "JSMITH"}, "not a persona"),
    ],
)
def test_a_persona_is_refused_with_what_to_fix(raw, why):
    with pytest.raises(ValueError, match=why):
        wbk.normalize_persona(raw, COLS, {"test-ca-mgr"})


def test_a_manager_must_be_another_persona():
    p = wbk.normalize_persona({"USERLOGIN": "test-x", "MGRLVL1": "test-ca-mgr"}, COLS, {"test-ca-mgr"})
    assert p["MGRLVL1"] == "test-ca-mgr"
