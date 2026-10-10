import io
import json
import time

import httpx
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from tdash.app import create_app
from tdash.config import Config
from tdash.workbook import diff, parse

ADMIN = {"X-Auth-Username": "ble", "X-Auth-Role": "admin"}
B = "/tools/testing-dash"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def cfg(tmp_path, **kw) -> Config:
    base = {"data_dir": str(tmp_path), "chatbot_url": "https://web.test", "chatbot_token": "svc-token",
                "judge_url": "https://j.test/v1", "judge_model": "judge-m", "git_commit": "deadbee"}  # fmt: skip
    base.update(kw)
    return Config(**base)


@pytest.fixture
def client(tmp_path, fakes):
    bot, judge = fakes
    app = create_app(
        cfg(tmp_path),
        env={},
        chatbot_transport=httpx.MockTransport(bot.handler),
        judge_transport=httpx.MockTransport(judge.handler),
    )
    with TestClient(app) as c:
        app.state.tdash["runner"].retry_delay = 0
        yield c


def imported(client, seed_bytes) -> dict:
    p = client.post(f"{B}/api/import/preview", headers=ADMIN, files={"file": ("set.xlsx", seed_bytes, XLSX)})
    assert p.status_code == 200, p.text
    a = client.post(f"{B}/api/import/apply", headers=ADMIN, json={"token": p.json()["token"]})
    assert a.status_code == 200, a.text
    return p.json()


def wait(client, run_id: int, timeout=20) -> dict:
    end = time.time() + timeout
    while time.time() < end:
        run = client.get(f"{B}/api/runs/{run_id}", headers=ADMIN).json()
        if run["status"] not in ("queued", "running"):
            return run
        time.sleep(0.05)
    raise AssertionError("run did not finish")


# ---------------------------------------------------------------- identity and the gate
def test_healthz_needs_no_identity(client):
    r = client.get("/healthz")
    assert r.json() == {"ok": True, "service": "testing-dash", "commit": "deadbee"}


@pytest.mark.parametrize(
    "headers",
    [{}, {"X-Auth-Username": "ble"}, {"X-Auth-Username": "ble", "X-Auth-Role": "user"}, {"X-Auth-Role": "admin"}],
)
def test_admins_only_everywhere(client, headers):
    for path in ("/", "/api/state", "/api/cases", "/api/runs", "/api/export/test-set"):
        assert client.get(B + path, headers=headers).status_code == 403
    assert client.post(f"{B}/api/runs", headers=headers, json={}).status_code == 403


def test_page_and_static(client):
    page = client.get(f"{B}/", headers=ADMIN).text
    assert f'data-base="{B}"' in page and f"{B}/static/app.js?v=" in page and "{{" not in page
    assert client.get(f"{B}/static/app.js").status_code == 200


def test_percent_encoded_identity(client):
    s = client.get(f"{B}/api/state", headers={"X-Auth-Username": "Jos%C3%A9", "X-Auth-Role": "admin"}).json()
    assert s["user"]["username"] == "josé"


def test_dev_identity_is_loopback_only(tmp_path):
    app = create_app(cfg(tmp_path, dev_identity="ble:admin"), env={})
    with TestClient(app) as c:  # TestClient's peer is "testclient", not loopback
        r = c.get(f"{B}/api/state")
        assert r.status_code == 403 and "loopback" in r.json()["error"]


def test_empty_base_path(tmp_path):
    app = create_app(cfg(tmp_path, base_path=""), env={})
    with TestClient(app) as c:
        assert c.get("/api/cases", headers=ADMIN).json() == {"cases": [], "version": 0}
        assert 'data-base=""' in c.get("/", headers=ADMIN).text


# ---------------------------------------------------------------- import / export
def test_import_round_trip(client, seed_bytes):
    pre = imported(client, seed_bytes)
    assert pre["cases"] == 75 and pre["personas"] == 12
    assert len(pre["diff"]["cases"]["added"]) == 75
    cases = client.get(f"{B}/api/cases", headers=ADMIN).json()
    assert len(cases["cases"]) == 75 and cases["version"] == 1
    out = client.get(f"{B}/api/export/test-set", headers=ADMIN)
    assert out.headers["content-type"] == XLSX and "hr-assist-test-set-v1.xlsx" in out.headers["content-disposition"]
    back = parse(out.content)
    d = diff({"cases": cases["cases"], "personas": back["personas"], "persona_columns": back["persona_columns"]}, back)
    assert d["cases"]["changed"] == [] and d["cases"]["added"] == d["cases"]["removed"] == []
    # importing the export again changes nothing
    again = client.post(f"{B}/api/import/preview", headers=ADMIN, files={"file": ("x.xlsx", out.content, XLSX)}).json()
    assert again["diff"]["cases"]["changed"] == [] and again["diff"]["personas"]["changed"] == []


def test_import_errors(client):
    r = client.post(f"{B}/api/import/preview", headers=ADMIN, files={"file": ("x.xlsx", b"junk", XLSX)})
    assert r.status_code == 422 and "not a readable" in r.json()["detail"]
    big = b"0" * (10 * 1024 * 1024 + 1)
    assert (
        client.post(f"{B}/api/import/preview", headers=ADMIN, files={"file": ("x.xlsx", big, XLSX)}).status_code == 413
    )
    assert client.post(f"{B}/api/import/apply", headers=ADMIN, json={"token": "nope"}).status_code == 410


def test_preview_tokens_are_single_use_and_personal(client, seed_bytes):
    tok = client.post(f"{B}/api/import/preview", headers=ADMIN, files={"file": ("s.xlsx", seed_bytes, XLSX)}).json()[
        "token"
    ]
    other = {"X-Auth-Username": "dustin", "X-Auth-Role": "admin"}
    assert client.post(f"{B}/api/import/apply", headers=other, json={"token": tok}).status_code == 403
    # someone else's attempt didn't burn it: the owner applies the same token
    assert client.post(f"{B}/api/import/apply", headers=ADMIN, json={"token": tok}).status_code == 200
    assert client.post(f"{B}/api/import/apply", headers=ADMIN, json={"token": tok}).status_code == 410


def test_roster_csv(client, seed_bytes):
    imported(client, seed_bytes)
    plain = client.get(f"{B}/api/export/test-roster.csv", headers=ADMIN)
    assert plain.headers["content-type"].startswith("text/csv") and "-10d" in plain.text
    assert "-10d" not in client.get(f"{B}/api/export/test-roster.csv?resolve_dates=true", headers=ADMIN).text


# ---------------------------------------------------------------- editing, versions
def test_edit_add_delete_and_versions(client, seed_bytes):
    imported(client, seed_bytes)
    c = client.get(f"{B}/api/cases", headers=ADMIN).json()["cases"][0]
    edited = client.put(f"{B}/api/cases/{c['id']}", headers=ADMIN, json={**c, "question": "Edited?", "review": None})
    assert edited.status_code == 200 and edited.json()["question"] == "Edited?" and edited.json()["_edited"]
    assert edited.json()["review"] is None and edited.json()["_updated_by"] == "ble"
    assert client.put(f"{B}/api/cases/{c['id']}", headers=ADMIN, json={**c, "id": "TC-X"}).status_code == 400
    assert client.put(f"{B}/api/cases/TC-NOPE", headers=ADMIN, json={"question": "q"}).status_code == 404
    new = client.post(
        f"{B}/api/cases", headers=ADMIN, json={"persona": "test-ca-on", "question": "New?", "expected": "x"}
    )
    assert new.json()["id"] == "TC-076" and new.json()["issues"] == []
    assert client.post(f"{B}/api/cases", headers=ADMIN, json={"id": "TC-076", "question": "q"}).status_code == 409
    assert client.post(f"{B}/api/cases", headers=ADMIN, json={"id": "TC 1/2", "question": "q"}).status_code == 400
    flagged = client.post(f"{B}/api/cases", headers=ADMIN, json={"persona": "BLE", "question": "q", "expected": "x"})
    assert flagged.status_code == 200 and "must be a test- login" in flagged.json()["issues"][0]  # kept, flagged
    assert client.delete(f"{B}/api/cases/TC-076", headers=ADMIN).status_code == 200
    assert client.delete(f"{B}/api/cases/TC-076", headers=ADMIN).status_code == 404
    vs = client.get(f"{B}/api/versions", headers=ADMIN).json()
    assert vs["current"] == 5 and [v["kind"] for v in vs["versions"]] == ["edit"] * 4 + ["import"]
    # a re-import would undo the in-app edit; the preview says so
    pre = client.post(f"{B}/api/import/preview", headers=ADMIN, files={"file": ("s.xlsx", seed_bytes, XLSX)}).json()
    assert c["id"] in pre["diff"]["cases"]["reverts_in_app_edits"]
    # roll back to the import
    assert client.post(f"{B}/api/versions/1/rollback", headers=ADMIN).json() == {"version": 6}
    after = {x["id"]: x for x in client.get(f"{B}/api/cases", headers=ADMIN).json()["cases"]}
    assert after[c["id"]]["question"] == c["question"] and "TC-077" not in after and not after[c["id"]]["_edited"]
    assert client.post(f"{B}/api/versions/99/rollback", headers=ADMIN).status_code == 404


def test_saving_an_unchanged_case_is_not_an_edit(client, seed_bytes):
    imported(client, seed_bytes)
    c = client.get(f"{B}/api/cases", headers=ADMIN).json()["cases"][0]
    body = {k: v for k, v in c.items() if not k.startswith("_")}
    again = client.put(f"{B}/api/cases/{c['id']}", headers=ADMIN, json=body).json()
    assert not again["_edited"]
    assert client.get(f"{B}/api/versions", headers=ADMIN).json()["current"] == 1  # no version written
    pre = client.post(f"{B}/api/import/preview", headers=ADMIN, files={"file": ("s.xlsx", seed_bytes, XLSX)}).json()
    assert pre["diff"]["cases"]["reverts_in_app_edits"] == []


# ---------------------------------------------------------------- runs
def test_run_lifecycle(client, seed_bytes, fakes):
    imported(client, seed_bytes)
    st = client.post(f"{B}/api/runs", headers=ADMIN, json={"scope": "ready", "concurrency": 4})
    assert st.status_code == 200, st.text
    run = wait(client, st.json()["id"])
    assert run["status"] == "done" and run["done"] == 65
    ov = run["overview"]
    assert ov["cases"] == 65 and ov["passed"] + ov["failed"] == 65 and ov["latency_p50_ms"] is not None
    assert run["info"]["chatbot_commit"] == "abc1234" and run["created_by"] == "ble"
    rs = client.get(f"{B}/api/runs/{run['id']}/results", headers=ADMIN).json()["results"]
    assert len(rs) == 65 and all(r["case"]["question"] for r in rs)
    x = client.get(f"{B}/api/runs/{run['id']}/export", headers=ADMIN)
    wb = load_workbook(io.BytesIO(x.content))
    assert wb.sheetnames == ["Results", "Summary"] and wb["Results"].max_row == 66
    assert client.post(f"{B}/api/runs/{run['id']}/cancel", headers=ADMIN).status_code == 409
    assert client.get(f"{B}/api/runs/999", headers=ADMIN).status_code == 404
    runs = client.get(f"{B}/api/runs", headers=ADMIN).json()["runs"]
    assert runs[0]["id"] == run["id"] and "audit" not in runs[0]
    # the second run is compared with the first
    second = wait(client, client.post(f"{B}/api/runs", headers=ADMIN, json={}).json()["id"])
    assert second["baseline_run"] == run["id"] and second["overview"]["has_baseline"]


def test_busy_blocks_runs_imports_and_rollbacks(client, seed_bytes, fakes):
    bot, _ = fakes
    bot.delay = 0.2
    imported(client, seed_bytes)
    rid = client.post(f"{B}/api/runs", headers=ADMIN, json={"concurrency": 1}).json()["id"]
    assert client.post(f"{B}/api/runs", headers=ADMIN, json={}).status_code == 409
    assert client.get(f"{B}/api/state", headers=ADMIN).json()["active_run"] == rid
    tok = client.post(f"{B}/api/import/preview", headers=ADMIN, files={"file": ("s.xlsx", seed_bytes, XLSX)}).json()[
        "token"
    ]
    assert client.post(f"{B}/api/import/apply", headers=ADMIN, json={"token": tok}).status_code == 409
    assert client.post(f"{B}/api/versions/1/rollback", headers=ADMIN).status_code == 409
    assert client.post(f"{B}/api/runs/{rid}/cancel", headers=ADMIN).status_code == 200
    assert wait(client, rid)["status"] == "cancelled"
    # the preview survived the 409 and still applies once the run is over
    assert client.post(f"{B}/api/import/apply", headers=ADMIN, json={"token": tok}).status_code == 200


def test_bad_run_requests(client, seed_bytes):
    assert client.post(f"{B}/api/runs", headers=ADMIN, json={}).json()["detail"].startswith("no runnable")
    imported(client, seed_bytes)
    assert client.post(f"{B}/api/runs", headers=ADMIN, json={"repeats": 7}).status_code == 400
    assert client.post(f"{B}/api/runs", headers=ADMIN, json={"concurrency": "lots"}).status_code == 400


def test_unconfigured_app(tmp_path, seed_bytes):
    app = create_app(Config(data_dir=str(tmp_path)), env={})
    with TestClient(app) as c:
        s = c.get(f"{B}/api/state", headers=ADMIN).json()
        assert not s["chatbot"]["ready"] and not s["judge"]["ready"]
        imported(c, seed_bytes)
        r = c.post(f"{B}/api/runs", headers=ADMIN, json={})
        assert r.status_code == 400 and "chatbot isn't configured" in r.json()["detail"]
        assert (
            c.post(f"{B}/api/playground", headers=ADMIN, json={"persona": "test-ca-on", "text": "hi"}).status_code
            == 503
        )


def test_secrets_resolve_at_startup(tmp_path, fakes, seed_bytes):
    bot, judge = fakes

    def pam(req):
        if "database" in req.url.path:
            return httpx.Response(200, json={"http_status": 200, "results": [{"field": [{"name": "TOK", "id": 1}]}]})
        return httpx.Response(200, json={"http_status": 200, "secret-id": 4738, "results": "svc-token"})

    env = {"CHATBOT_TOKEN": "SECRET:TOK", "FORTIPAM_TOKEN": "t", "FORTIPAM_SECRET_ID": "4738"}
    app = create_app(
        cfg(tmp_path, chatbot_token="SECRET:TOK"),
        env=env,
        chatbot_transport=httpx.MockTransport(bot.handler),  # asserts the resolved Bearer token
        judge_transport=httpx.MockTransport(judge.handler),
        pam_transport=httpx.MockTransport(pam),
    )
    with TestClient(app) as c:
        imported(c, seed_bytes)
        r = c.post(f"{B}/api/playground", headers=ADMIN, json={"persona": "test-ca-on", "text": "hi"})
        assert r.status_code == 200


def test_secret_failure_stops_the_boot(tmp_path):
    app = create_app(cfg(tmp_path), env={"CHATBOT_TOKEN": "SECRET:X"})
    with pytest.raises(Exception, match="FORTIPAM"), TestClient(app):
        pass


# ---------------------------------------------------------------- playground
def test_playground(client, seed_bytes, fakes):
    bot, _ = fakes
    imported(client, seed_bytes)
    p = lambda body: client.post(f"{B}/api/playground", headers=ADMIN, json=body)  # noqa: E731
    assert p({"persona": "BLE", "text": "hi"}).status_code == 400
    assert p({"persona": "test-nobody", "text": "hi"}).status_code == 400
    assert p({"persona": "test-ca-on", "text": "  "}).status_code == 400
    assert p({"persona": "test-ca-on", "text": "hi", "history": [{"role": "system", "text": "x"}]}).status_code == 400
    assert (
        p({"persona": "test-ca-on", "text": "hi", "history": [{"role": "user", "text": "x"}] * 21}).status_code == 400
    )
    hist = [{"role": "user", "text": "first"}, {"role": "assistant", "text": "reply"}]
    ok = p({"persona": "test-ca-qc", "text": "and then?", "history": hist}).json()
    assert ok["answer"].startswith("Here is the answer") and ok["history_supported"] is False
    assert bot.calls[-1] == {"sender": "test-ca-qc", "text": "and then?", "history": hist}
    bot.fail["boom"] = [500]
    r = p({"persona": "test-ca-on", "text": "boom"})
    assert r.status_code == 502 and "HTTP 500" in r.json()["detail"]


def test_restart_interrupts_a_running_run(tmp_path, fakes, seed_bytes):
    bot, judge = fakes
    bot.delay = 0.3
    t = {"chatbot_transport": httpx.MockTransport(bot.handler), "judge_transport": httpx.MockTransport(judge.handler)}
    app = create_app(cfg(tmp_path), env={}, **t)
    with TestClient(app) as c:
        imported(c, seed_bytes)
        rid = c.post(f"{B}/api/runs", headers=ADMIN, json={"concurrency": 1}).json()["id"]
        time.sleep(0.4)
    # shutdown recorded it; a fresh app on the same data sees the same thing
    with TestClient(create_app(cfg(tmp_path), env={}, **t)) as c:
        run = c.get(f"{B}/api/runs/{rid}", headers=ADMIN).json()
        assert run["status"] == "interrupted" and 0 < run["done"] < run["total"]
        assert json.loads(json.dumps(run))  # serialisable end to end
